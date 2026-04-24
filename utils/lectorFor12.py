import openpyxl
from tkinter import filedialog
import tkinter as tk
import random
from datetime import datetime
from utils import *
# --- Lógica de Lectura ---

def tipoFor12d():
    try:
        root = tk.Tk()
        root.withdraw()
        nombre_libro = filedialog.askopenfilename()
        if not nombre_libro: return [False, None]
        
        libro = openpyxl.load_workbook(nombre_libro, data_only=True)
        hoja = libro['Plan de trabajo']

        numCelda = 1
        for fila in hoja.iter_rows():
            if numCelda >= 7:
                # Detección de formato basado en la primera celda
                if not isinstance(fila[0].value, int):
                    return [True, libro]
                else:
                    return [False, libro]
            numCelda += 1
    except:
        return [False, None]

def for12Data(archivo):
    """Lectura estándar del FOR12 sin distinción de previas/ventana"""
    nameCampos = []
    rows = []

    try:
        libro = archivo
        hoja = libro['Plan de trabajo']
        numFila = 1
        campos_crudos = []

        for fila in hoja.iter_rows():
            if numFila == 5:
                campos_crudos = [celda.value for celda in fila]
                nameCampos = limpiar_campos(campos_crudos)
            
            elif numFila >= 7:
                if not isinstance(fila[0].value, int):
                    numFila += 1
                    continue
                
                diccionario = {}
                valores_fila = [celda.value for i, celda in enumerate(fila) 
                               if i < len(campos_crudos) and campos_crudos[i] is not None]
                
                for i, valor in enumerate(valores_fila):
                    if i < len(nameCampos):
                        diccionario[nameCampos[i]] = valor
                
                if diccionario:
                    rows.append(diccionario)
            numFila += 1

        # Uso de la función modularizada
        datosGenerales = obtener_datos_generales(libro)

        return [eliminar_vacios(rows), datosGenerales]
    except Exception as e:
        print(f"Error inesperado al leer el archivo (for12Data): {e}")
        return [[], {}]

def conActividadesPrevias(archivo):
    """Lectura dividida: detecta Previas y luego llama a Ventana"""
    nameCampos = []
    rows_previas = []
    libro = archivo
    hoja = libro['Plan de trabajo']
    
    numFilaContador = 1
    breakear = False
    contBool = 0
    campos_originales_con_none = []

    for fila in hoja.iter_rows():
        if breakear:
            break
        
        if numFilaContador == 5:
            campos_originales_con_none = [celda.value for celda in fila]
            nameCampos = limpiar_campos(campos_originales_con_none)
            
        elif numFilaContador >= 8:
            # Validación de fin de bloque de previas
            try:
                if fila[0].value is None: raise ValueError
                int(fila[0].value)
            except:
                breakear = True
                contBool += 1
                break
            
            diccionario = {}
            valores_fila = [celda.value for i, celda in enumerate(fila) 
                           if i < len(campos_originales_con_none) and campos_originales_con_none[i] is not None]
            
            for i, valor in enumerate(valores_fila):
                if i < len(nameCampos):
                    diccionario[nameCampos[i]] = valor
            
            if diccionario:
                rows_previas.append(diccionario)
        
        numFilaContador += 1

    ventanaActividade = actividadesVentana(numFilaContador, libro, nameCampos, campos_originales_con_none)
    
    # Uso de la función modularizada
    datosGenerales = obtener_datos_generales(libro)

    return [ventanaActividade, eliminar_vacios(rows_previas), datosGenerales]

def actividadesVentana(numFilaInicio, archivo, nameCamposLimpios, camposOriginales):
    """Lee las actividades principales después de las previas"""
    rows_ventana = []
    hoja = archivo['Plan de trabajo']
    
    for fila in hoja.iter_rows(min_row=numFilaInicio):
        if isinstance(fila[0].value, int):
            diccionario = {}
            valores_fila = [celda.value for i, celda in enumerate(fila) 
                           if i < len(camposOriginales) and camposOriginales[i] is not None]
            
            for i, valor in enumerate(valores_fila):
                if i < len(nameCamposLimpios):
                    diccionario[nameCamposLimpios[i]] = valor
            
            if diccionario:
                rows_ventana.append(diccionario)

    return eliminar_vacios(rows_ventana)

# --- Mensajería ---

def mensajeFinalRandom():
    messages = [
        "Agradezco sus comentarios y solicito el envío de las evidencias.",
        "Quedo atento a sus comentarios y a la recepción de las evidencias.",
        "Estaré pendiente de sus comentarios y del envío de las evidencias.",
        "Agradezco de antemano sus comentarios y el envío de las evidencias."
    ]
    return random.choice(messages)

def mensajeContinuidadRandom():
    messages = [
        "dando continuidad con la For12D,",
        "continuando con el trabajo,",
        "siguiendo con las actividades,",
        "prosiguiendo con lo estipulado en la For12D,"
    ]
    return random.choice(messages)

def mensajeFinActividades():
    messages = [
        "Se confirma la ejecución exitosa de todas las actividades especificadas en la For12d: ",
        "Se notifica que todas las actividades detalladas en la For12d han sido realizadas de manera efectiva: ",
        "Se valida la correcta ejecución de todas las actividades descritas en la For12d: "
    ]
    return random.choice(messages)